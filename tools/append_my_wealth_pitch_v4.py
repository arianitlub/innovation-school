from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
src = ROOT / 'docs/2-challenge-wealth-advice/presentations/my-wealth-5min-pitch-v3.pptx'
out = ROOT / 'docs/2-challenge-wealth-advice/presentations/my-wealth-5min-pitch-v4-business-case.pptx'
focus_img = ROOT / 'tools/assets/business-case-focus.jpg'
ingredients_img = ROOT / 'tools/assets/business-case-ingredients.jpg'

prs = Presentation(src)
SW = prs.slide_width
Y='FFD500'; Y2='FFF0A8'; INK='171717'; MUTED='6E6A62'; BG='FFFDF3'; W='FFFFFF'; LINE='E6DFC7'; GREEN='2E7D5A'; PALE='FFF8D8'

def rgb(h): return RGBColor.from_string(h)

def add_text(sl, txt, x,y,w,h, size=18,bold=False,color=INK,align='left',font='Liberation Sans'):
    tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame; tf.clear(); tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.text=txt; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
    p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(align,PP_ALIGN.LEFT)
    return tb

def rect(sl,x,y,w,h,fill=W,line=LINE,radius=True,lw=1):
    typ=MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh=sl.shapes.add_shape(typ,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    sh.line.color.rgb=rgb(line); sh.line.width=Pt(lw)
    return sh

def pill(sl,txt,x,y,w,fill=Y,color=INK):
    rect(sl,x,y,w,0.34,fill,fill,True,0.5)
    add_text(sl,txt,x+0.05,y,w-0.1,0.34,9,True,color,'center')

def head(sl,k,title,sub=None):
    pill(sl,k,0.65,0.42,1.65)
    add_text(sl,title,0.65,0.92,12,0.66,28,True)
    if sub: add_text(sl,sub,0.65,1.58,11.9,0.42,13,False,MUTED)

def add_footer(sl,num):
    add_text(sl,'MY WEALTH',11.35,7.03,1.45,0.18,8,True,'77736A','right')
    add_text(sl,str(num),12.85,7.03,0.25,0.18,8,False,'8E8A82','right')

def add_image_crop(sl,path,x,y,w,h,rotate=0):
    im=Image.open(path).convert('RGB')
    target=w/h; ar=im.width/im.height
    if ar>target:
        nw=int(im.height*target); left=(im.width-nw)//2; im=im.crop((left,0,left+nw,im.height))
    else:
        nh=int(im.width/target); top=(im.height-nh)//2; im=im.crop((0,top,im.width,top+nh))
    temp=ROOT / 'tools/assets/_pitch_v4_crop.jpg'
    im.save(temp,quality=88)
    pic=sl.shapes.add_picture(str(temp),Inches(x),Inches(y),Inches(w),Inches(h)); pic.rotation=rotate
    return pic

blank=prs.slide_layouts[0]

# Slide 8 - where we are
sl=prs.slides.add_slide(blank)
sl.background.fill.solid(); sl.background.fill.fore_color.rgb=rgb(BG)
r=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,SW,Inches(0.08)); r.fill.solid(); r.fill.fore_color.rgb=rgb(Y); r.line.color.rgb=rgb(Y)
head(sl,'WHERE WE ARE','Stage 5: turn a good idea into an investable decision','We have a problem, evidence, concept and prototype. Now we prove numbers, delivery and next-step gates.')
x0=0.88; y=2.55; total_w=7.15
rail=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x0+0.25),Inches(y+0.20),Inches(total_w-0.5),Inches(0.05)); rail.fill.solid(); rail.fill.fore_color.rgb=rgb('D8D1B8'); rail.line.color.rgb=rgb('D8D1B8')
stages=[('1','UNDERSTAND','Challenge + scope','done'),('2','EMPATHIZE','Customers + experts','done'),('3','IDEATE','70 ideas -> My Wealth','done'),('4','PROTOTYPE','T1 feedback -> iterate','partial'),('5','BUSINESS CASE','Numbers + action plan','now'),('6','PILOT','Real cohort + KPIs','next')]
step=total_w/5.6
for i,(n,name,desc,status) in enumerate(stages):
    xx=x0+i*step
    fill=Y if status in ('done','now') else (Y2 if status=='partial' else 'F0ECE0')
    line=Y if status!='next' else 'CFC8AF'
    c=sl.shapes.add_shape(MSO_SHAPE.OVAL,Inches(xx),Inches(y),Inches(0.48),Inches(0.48)); c.fill.solid(); c.fill.fore_color.rgb=rgb(fill); c.line.color.rgb=rgb(line)
    add_text(sl,'✓' if status=='done' else ('●' if status=='now' else n),xx,y+0.01,0.48,0.46,13,True,INK,'center')
    add_text(sl,name,xx-0.28,y+0.72,1.05,0.28,9,True,INK,'center')
    add_text(sl,desc,xx-0.43,y+1.04,1.35,0.48,8.5,False,MUTED,'center')
    if status=='now': pill(sl,'WE ARE HERE',xx-0.22,y-0.54,0.92,Y,INK)
rect(sl,0.78,4.66,7.40,1.52,W,LINE)
add_text(sl,'What is already strong',1.02,4.88,2.25,0.28,14,True)
for j,txt in enumerate(['Problem & HMW are clear','Early customer + expert evidence','Prototype makes the concept tangible']):
    add_text(sl,'✓',1.02,5.29+j*0.27,0.22,0.22,10,True,GREEN,'center'); add_text(sl,txt,1.30,5.27+j*0.27,2.55,0.24,10,False)
add_text(sl,'What Stage 5 must prove',4.18,4.88,2.4,0.28,14,True)
for j,txt in enumerate(['How many eligible customers?','What uplift / margin is realistic?','What team, cost and controls are needed?']):
    add_text(sl,'->',4.18,5.29+j*0.27,0.25,0.22,10,True,INK,'center'); add_text(sl,txt,4.46,5.27+j*0.27,3.1,0.24,10,False)
rect(sl,8.50,2.15,4.08,4.36,W,'D9D2BA')
add_image_crop(sl,focus_img,8.72,2.39,3.64,2.62,-1)
pill(sl,'RBI STAGE 5',8.86,5.23,1.12,Y2,INK)
add_text(sl,'Numbers',8.86,5.67,0.95,0.24,12,True); add_text(sl,'Evidence & arguments',9.82,5.67,1.55,0.24,12,True); add_text(sl,'Action plan',11.38,5.67,0.78,0.24,12,True)
add_text(sl,'Round 1 feedback: clearer Wealth Insights + notifications. Round 2: 5-8 people next.',8.86,6.00,3.20,0.42,9.5,False,MUTED)
add_footer(sl,8)

# Slide 9 - team and action plan
sl2=prs.slides.add_slide(blank)
sl2.background.fill.solid(); sl2.background.fill.fore_color.rgb=rgb(BG)
r=sl2.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,SW,Inches(0.08)); r.fill.solid(); r.fill.fore_color.rgb=rgb(Y); r.line.color.rgb=rgb(Y)
head(sl2,'ACTION PLAN','Make it real: one squad, six months, clear stop/go gates','A banking-grade pilot needs a small core squad plus shared specialists from across the bank.')
rect(sl2,0.72,2.15,4.05,4.62,INK,INK); pill(sl2,'TEAM',0.98,2.40,0.72,Y,INK)
add_text(sl2,'~14 people involved',0.98,2.94,3.18,0.38,23,True,W); add_text(sl2,'~8-9 blended FTE during the pilot',0.98,3.32,3.18,0.30,11,False,'D6D6D6')
add_text(sl2,'Core squad',0.98,3.86,1.45,0.27,13,True,Y)
core=[('1','Product Owner'),('1','UX / Product Designer'),('2','Mobile Engineers'),('1','Backend / API Engineer'),('1','Data / AI Engineer'),('1','QA / Automation'),('0.5','Tech Lead / Architect')]
for i,(n,role) in enumerate(core):
    yy=4.18+i*0.32; add_text(sl2,n,1.00,yy,0.35,0.23,10,True,Y,'center'); add_text(sl2,role,1.42,yy,2.70,0.23,10,False,W)
add_text(sl2,'Shared specialists (part-time)',0.98,6.46,2.4,0.25,11,True,'CFCFCF')
pill(sl2,'SHARED BANK EXPERTISE',5.18,2.15,1.75,Y2,INK)
shared=[('INVESTMENT','Advisory / fund SME'),('COMPLIANCE','MiFID + legal'),('SECURITY','Privacy + cyber'),('RISK','AI / model governance'),('DATA','Analytics + segmentation'),('CHANGE','Ops + training')]
for i,(a,b) in enumerate(shared):
    xx=5.18+(i%3)*2.40; yy=2.67+(i//3)*0.86; rect(sl2,xx,yy,2.18,0.68,W,LINE); add_text(sl2,a,xx+0.12,yy+0.08,0.92,0.19,9,True); add_text(sl2,b,xx+0.12,yy+0.31,1.88,0.23,9,False,MUTED)
add_text(sl2,'6-month controlled pilot',5.18,4.43,3.10,0.32,18,True)
steps=[('W1-4','PROVE','Cohort, KPIs, round-2 tests, Compliance/MiFID/GDPR'),('W5-12','BUILD','Life Profile, Learn, Playground, mobile + backend'),('W13-18','CONNECT','AI explainer guardrails, adviser handover, analytics, security'),('W19-24','PILOT','Controlled launch, measure activation, AUM, satisfaction, capacity')]
for i,(wk,name,desc) in enumerate(steps):
    yy=4.92+i*0.49; pill(sl2,wk,5.18,yy,0.72,Y if i<3 else Y2,INK); add_text(sl2,name,6.03,yy,0.78,0.29,11,True); add_text(sl2,desc,6.84,yy,4.48,0.29,8.6,False,MUTED)
rect(sl2,5.18,6.88,7.30,0.42,PALE,'E4D16A'); add_text(sl2,'GATE: scale only if customer value + economics + controls are proven',5.36,6.95,6.92,0.24,11,True,INK,'center')
add_image_crop(sl2,ingredients_img,11.55,4.62,0.92,1.76,1)
pill(sl2,'€350k PILOT ENVELOPE',9.58,4.43,1.95,Y,INK); add_text(sl2,'Working assumption; validate with RBI delivery owners.',9.58,4.77,1.90,0.38,8.5,False,MUTED,'right')
add_footer(sl2,9)

# Put the two new slides before the original investment / upside / ask slides.
sldIdLst=prs.slides._sldIdLst
new1=sldIdLst[-2]; new2=sldIdLst[-1]
sldIdLst.remove(new1); sldIdLst.remove(new2); sldIdLst.insert(7,new1); sldIdLst.insert(8,new2)
prs.save(out)
